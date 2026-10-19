#!/usr/bin/env python3
"""
SEAP (e-licitatie.ro) scraper: recent construction tenders above a RON threshold.
Outputs:  docs/tenders.xlsx   and   docs/index.html   (static site)

Usage:
    python scraper.py            # normal run
    python scraper.py --debug    # print raw first item so you can fix field names
"""
import sys, json, html, datetime as dt
from pathlib import Path
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------- settings you can change ----------------
MIN_VALUE_RON = 30_000_000
DAYS_BACK     = 30          # how far back to look
CPV_PREFIX    = "45"        # CPV division 45 = construction works
PAGE_SIZE     = 50
MAX_PAGES     = 40
# ----------------------------------------------------------

BASE = "https://www.e-licitatie.ro"
API  = f"{BASE}/api-pub/NoticeCommon/GetCNoticeList/"
OUT  = Path("docs"); OUT.mkdir(exist_ok=True)
DEBUG = "--debug" in sys.argv

HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/124 Safari/537.36",
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json",
    "Origin": BASE,
    "Referer": f"{BASE}/pub/notices/contract-notices/list/0/0",
    "Accept-Language": "ro-RO,ro;q=0.9,en;q=0.8",
}


def payload(page_index, start_date):
    return {
        "sysNoticeTypeIds": [],
        "sysProcedureStateId": None,
        "contractingAuthorityId": None,
        "winnerId": None,
        "cpvCodeId": None,
        "contractTypeId": None,
        "startPublicationDate": start_date,
        "endPublicationDate": None,
        "estimatedValueRonFrom": MIN_VALUE_RON,   # server-side filter (ignored if name differs)
        "estimatedValueRonTo": None,
        "sortProperties": [],
        "sortOrder": [],
        "pageSize": PAGE_SIZE,
        "pageIndex": page_index,
        "hasUnansweredQuestions": False,
    }


def first(d, *keys, default=""):
    for k in keys:
        if d.get(k) not in (None, ""):
            return d[k]
    return default


def fetch_all():
    s = requests.Session()
    s.headers.update(HEADERS)
    s.get(f"{BASE}/pub/notices/contract-notices/list/0/0", timeout=30)  # pick up cookies
    start = (dt.datetime.utcnow() - dt.timedelta(days=DAYS_BACK)).strftime("%Y-%m-%dT00:00:00.000Z")
    items = []
    for p in range(MAX_PAGES):
        r = s.post(API, data=json.dumps(payload(p, start)), timeout=60)
        r.raise_for_status()
        data = r.json()
        page = data.get("items") or data.get("Items") or []
        if DEBUG and p == 0:
            print("TOTAL:", data.get("total"))
            print(json.dumps(page[:1], indent=2, ensure_ascii=False))
        items += page
        total = data.get("total") or data.get("Total") or 0
        if not page or len(items) >= total:
            break
    return items


def normalise(it):
    cpv = str(first(it, "cpvCodeAndName", "cpvCode"))
    value = first(it, "ronContractValue", "estimatedValueRon", "ronValue", default=0) or 0
    nid = first(it, "cNoticeId", "noticeId", "id")
    return {
        "published": str(first(it, "noticeStateDate", "publicationDate", "noticeDate"))[:10],
        "deadline": str(first(it, "tenderReceiptDeadline", "receiptDeadline"))[:16].replace("T", " "),
        "title": first(it, "contractTitle", "title"),
        "authority": first(it, "contractingAuthorityNameAndFiscalNumber", "contractingAuthorityName"),
        "value": float(value),
        "cpv": cpv,
        "type": first(it, "sysAcquisitionContractType", "contractType"),
        "procedure": first(it, "sysProcedureType", "procedureType"),
        "notice_no": first(it, "noticeNo", "noticeNumber"),
        "url": f"{BASE}/pub/notices/c-notice/v2/view/{nid}" if nid else "",
    }


def is_construction(r):
    return r["cpv"].strip().startswith(CPV_PREFIX) or "lucr" in str(r["type"]).lower()


def write_excel(rows):
    wb = Workbook(); ws = wb.active; ws.title = "Construction tenders"
    cols = [("Published", "published", 12), ("Deadline", "deadline", 17), ("Title", "title", 60),
            ("Authority", "authority", 45), ("Value (RON)", "value", 18), ("CPV", "cpv", 40),
            ("Procedure", "procedure", 25), ("Notice no.", "notice_no", 14), ("Link", "url", 40)]
    ws.append([c[0] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in rows:
        ws.append([r[k] for _, k, _ in cols])
        ws.cell(ws.max_row, 5).number_format = "#,##0"
        if r["url"]:
            ws.cell(ws.max_row, 9).hyperlink = r["url"]
    for i, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / "tenders.xlsx")


def write_html(rows):
    e = html.escape
    body = "".join(
        f"<tr><td>{e(r['published'])}</td><td>{e(r['deadline'])}</td>"
        f"<td><a href='{e(r['url'])}' target='_blank'>{e(str(r['title']))}</a></td>"
        f"<td>{e(str(r['authority']))}</td><td data-v='{r['value']:.0f}' class='n'>{r['value']:,.0f}</td>"
        f"<td>{e(r['cpv'])}</td></tr>" for r in rows)
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Construction tenders &gt; {MIN_VALUE_RON/1e6:.0f}M RON</title>
<style>
:root{{--bg:#fff;--fg:#1a1a1a;--mut:#666;--line:#e3e3e3;--acc:#1f3a5f;--hov:#f4f7fb}}
@media(prefers-color-scheme:dark){{:root{{--bg:#16181c;--fg:#eee;--mut:#9aa;--line:#2c3036;--acc:#8fb4ea;--hov:#1e2228}}}}
body{{font:15px system-ui,sans-serif;background:var(--bg);color:var(--fg);margin:0;padding:24px}}
h1{{margin:0 0 4px;font-size:22px}} .m{{color:var(--mut);margin-bottom:16px}}
input{{padding:8px 10px;width:100%;max-width:360px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);margin-bottom:12px}}
.w{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;min-width:900px}}
th,td{{padding:8px 10px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}}
th{{cursor:pointer;color:var(--acc);position:sticky;top:0;background:var(--bg)}} tr:hover td{{background:var(--hov)}}
a{{color:var(--acc)}} .n{{text-align:right;white-space:nowrap}}
</style></head><body>
<h1>Construction tenders above {MIN_VALUE_RON:,.0f} RON</h1>
<div class="m">{len(rows)} notices from the last {DAYS_BACK} days · updated {now} · <a href="tenders.xlsx">Download Excel</a></div>
<input id="q" placeholder="Filter…" oninput="f()">
<div class="w"><table id="t"><thead><tr><th>Published</th><th>Deadline</th><th>Title</th><th>Authority</th><th>Value (RON)</th><th>CPV</th></tr></thead>
<tbody>{body}</tbody></table></div>
<script>
function f(){{const q=document.getElementById('q').value.toLowerCase();
document.querySelectorAll('#t tbody tr').forEach(r=>r.style.display=r.textContent.toLowerCase().includes(q)?'':'none')}}
document.querySelectorAll('th').forEach((h,i)=>h.onclick=()=>{{const b=document.querySelector('#t tbody');
const rs=[...b.rows];const d=h.dataset.d=h.dataset.d==='1'?'-1':'1';
rs.sort((a,c)=>{{const x=a.cells[i].dataset.v??a.cells[i].textContent,y=c.cells[i].dataset.v??c.cells[i].textContent;
return (isNaN(x)||isNaN(y)?x.localeCompare(y):x-y)*d}});rs.forEach(r=>b.appendChild(r))}});
</script></body></html>"""
    (OUT / "index.html").write_text(page, encoding="utf-8")


def main():
    raw = fetch_all()
    rows = [normalise(i) for i in raw]
    rows = [r for r in rows if r["value"] >= MIN_VALUE_RON and is_construction(r)]
    rows.sort(key=lambda r: r["published"], reverse=True)
    write_excel(rows); write_html(rows)
    print(f"Fetched {len(raw)} notices, kept {len(rows)} -> docs/tenders.xlsx, docs/index.html")


if __name__ == "__main__":
    main()
